import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Colors
BG_COLOR = RGBColor(11, 11, 11) # #0B0B0B
CARD_COLOR = RGBColor(26, 26, 26) # Dark Charcoal
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_GRAY = RGBColor(180, 180, 180)
ACCENT_BLUE = RGBColor(0, 255, 255) # Cyan Glow
ACCENT_PURPLE = RGBColor(157, 78, 221)

prs = Presentation()
prs.slide_width = Inches(13.333) # 16:9 widescreen
prs.slide_height = Inches(7.5)

# Function to set slide background
def set_bg(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_blank_slide():
    slide_layout = prs.slide_layouts[6] # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_bg(slide)
    return slide

def add_text(slide, text, left, top, width, height, size, bold=False, color=TEXT_WHITE, align=PP_ALIGN.LEFT, font_name="Inter"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = size
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return txBox

def add_card(slide, left, top, width, height, text_title, text_body):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_COLOR
    shape.line.color.rgb = RGBColor(50, 50, 50) # Subtle border
    
    add_text(slide, text_title, left + Inches(0.3), top + Inches(0.3), width - Inches(0.6), Inches(0.5), Pt(20), bold=True)
    if text_body:
        add_text(slide, text_body, left + Inches(0.3), top + Inches(0.8), width - Inches(0.6), height - Inches(1.1), Pt(14), color=TEXT_GRAY)

# Slide 1: Hero
slide1 = add_blank_slide()
add_text(slide1, "PRANESH M S", Inches(1), Inches(2.5), Inches(11.33), Inches(1), Pt(64), bold=True, align=PP_ALIGN.CENTER)
add_text(slide1, "A Curious Learner | Future ML Engineer | Builder", Inches(1), Inches(4), Inches(11.33), Inches(1), Pt(24), color=ACCENT_BLUE, align=PP_ALIGN.CENTER)

# Slide 2: Bento Grid (About Me)
slide2 = add_blank_slide()
add_text(slide2, "About Me", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
add_card(slide2, Inches(1), Inches(1.5), Inches(3.5), Inches(5), "📍 Origins", "Raised in Salem\n\nSchooling at SKV CBSE School, Namakkal.")
add_card(slide2, Inches(4.8), Inches(1.5), Inches(7.5), Inches(2.3), "🎓 Current Chapter", "B.Tech CSE (2nd Year) at SRM Institute of Science and Technology.")
add_card(slide2, Inches(4.8), Inches(4.1), Inches(7.5), Inches(2.4), "🧠 Mindset", "Not an expert yet, but a relentless learner who adapts continuously.")

# Slide 3: Beyond Academics (Cards)
slide3 = add_blank_slide()
add_text(slide3, "Beyond the Screen", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
add_card(slide3, Inches(1), Inches(2), Inches(3.5), Inches(4), "🏏 Cricket", "Right-Arm Medium Fast Bowler. The pitch is my reset button.")
add_card(slide3, Inches(4.8), Inches(2), Inches(3.5), Inches(4), "🏍️ Riding", "Bike Enthusiast. Finding clarity on the open road.")
add_card(slide3, Inches(8.6), Inches(2), Inches(3.5), Inches(4), "🎬 Cinema", "Movie & Web Series Buff. Finding inspiration in storytelling.")

# Slide 4: Emotional Storytelling (Learner's Mindset)
slide4 = add_blank_slide()
add_text(slide4, '"', Inches(1), Inches(1), Inches(11.33), Inches(1), Pt(100), color=ACCENT_PURPLE, align=PP_ALIGN.CENTER)
add_text(slide4, "I believe learning is a lifelong journey, \nand every project is an opportunity to grow.", Inches(1), Inches(2.5), Inches(11.33), Inches(2), Pt(36), bold=True, align=PP_ALIGN.CENTER)
add_text(slide4, "Being willing to learn is more valuable than pretending to know everything.", Inches(1), Inches(5), Inches(11.33), Inches(1), Pt(20), color=TEXT_GRAY, align=PP_ALIGN.CENTER)

# Slide 5: Core Strengths (Chips)
slide5 = add_blank_slide()
add_text(slide5, "Core Strengths", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
strengths = ["Problem Solving", "Leadership", "Hard Working", "Self-Learning", "Curiosity", "Adaptability", "Perfection"]
# Draw chips
x, y = Inches(1), Inches(2.5)
for s in strengths:
    w = Inches(3)
    h = Inches(1)
    shape = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_COLOR
    shape.line.color.rgb = RGBColor(50, 50, 50)
    add_text(slide5, s, x, y+Inches(0.2), w, Inches(0.6), Pt(18), align=PP_ALIGN.CENTER)
    x += Inches(3.5)
    if x > Inches(10):
        x = Inches(1)
        y += Inches(1.5)

# Slide 6: Current Focus (Bento)
slide6 = add_blank_slide()
add_text(slide6, "Current Focus", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
add_card(slide6, Inches(1), Inches(1.5), Inches(5.5), Inches(2.5), "🤖 Machine Learning", "Learning ML from fundamentals to build intelligent products.")
add_card(slide6, Inches(6.8), Inches(1.5), Inches(5.5), Inches(2.5), "📱 App Development", "Using Flutter to convert web platforms into seamless mobile experiences.")
add_card(slide6, Inches(1), Inches(4.3), Inches(11.3), Inches(2.2), "⚙️ DevOps & AI Tools", "Deploying complete applications to the internet and utilizing AI-assisted development workflows to accelerate building.")

# Slide 7: Tech Philosophy
slide7 = add_blank_slide()
add_text(slide7, "The Philosophy", Inches(1), Inches(2.5), Inches(4), Inches(2), Pt(40), bold=True)
add_text(slide7, "Software engineering has no limits. \nTechnologies evolve daily.\n\nMaintain a broad mind across domains while building deep specialization in Machine Learning.", Inches(5), Inches(2.5), Inches(7), Inches(4), Pt(24), color=TEXT_GRAY)
line = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(2), Inches(0.05), Inches(4))
line.fill.solid()
line.fill.fore_color.rgb = ACCENT_BLUE

# Slide 8: Real-World Exposure
slide8 = add_blank_slide()
add_text(slide8, "Real-World Exposure", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
add_card(slide8, Inches(1), Inches(2), Inches(5.5), Inches(4), "AI & ML Intern", "Integrated AI into real-world products.\nMastered modern AI-assisted workflows.\nCollaborated with experienced developers.")
add_card(slide8, Inches(6.8), Inches(2), Inches(5.5), Inches(4), "Cybersecurity Intern", "Learned secure development practices.\nUnderstood vulnerability assessments.\nExplored secure application architectures.")

# Slide 9: Achievements
slide9 = add_blank_slide()
add_text(slide9, "Milestones", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
add_card(slide9, Inches(1), Inches(1.5), Inches(11.33), Inches(1.5), "🏆 Hack for Bharat - Winner", "Recognized for building impact-driven solutions for real-world problems.")
add_card(slide9, Inches(1), Inches(3.3), Inches(11.33), Inches(1.5), "💼 SWE at ML Startup", "Selected as a Part-Time Software Developer to gain practical industry experience.")
add_card(slide9, Inches(1), Inches(5.1), Inches(11.33), Inches(1.5), "🚀 Hybrid Startup Vision", "Initiated the roadmap for a startup combining both Products and Services.")

# Slide 10: 5-Year Vision
slide10 = add_blank_slide()
add_text(slide10, "Where Am I Going?", Inches(1), Inches(1), Inches(11.33), Inches(1), Pt(40), bold=True, align=PP_ALIGN.CENTER)
add_text(slide10, "Five years from now, I see myself as a \nMachine Learning Engineer.", Inches(1), Inches(3), Inches(11.33), Inches(2), Pt(32), color=ACCENT_BLUE, align=PP_ALIGN.CENTER)
add_text(slide10, "Working on intelligent systems that solve meaningful problems.", Inches(1), Inches(4.5), Inches(11.33), Inches(1), Pt(20), color=TEXT_GRAY, align=PP_ALIGN.CENTER)

# Slide 11: Roadmap - Year 1
slide11 = add_blank_slide()
add_text(slide11, "Roadmap: Year 1", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
add_card(slide11, Inches(1), Inches(2), Inches(11.33), Inches(4), "The Entrepreneurial Spark", "Launch a small startup initiative based on a hybrid model (Products + Services).\n\nKeep it as a part-time engagement to build the foundation while finishing studies.")

# Slide 12: Roadmap - Years 3 to 5
slide12 = add_blank_slide()
add_text(slide12, "Roadmap: Years 3-5", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
add_card(slide12, Inches(1), Inches(2), Inches(11.33), Inches(4), "The Industry Phase", "Work as a Machine Learning Engineer in a leading technology company.\n\nLearn from experienced professionals, understand scalable product design, and gain deep industry experience.")

# Slide 13: Roadmap - Years 5 to 10
slide13 = add_blank_slide()
add_text(slide13, "Roadmap: Years 5-10", Inches(1), Inches(0.5), Inches(11.33), Inches(1), Pt(40), bold=True)
add_card(slide13, Inches(1), Inches(2), Inches(11.33), Inches(4), "The Leap", "Return to my startup full-time with the knowledge, experience, and leadership gained from the industry.\n\nScale it into a company that builds impactful products while solving real-world problems.")

# Slide 14: The "Why"
slide14 = add_blank_slide()
add_text(slide14, "Everyone should have a cause \nbehind what they do.", Inches(1), Inches(2), Inches(11.33), Inches(2), Pt(40), bold=True, align=PP_ALIGN.CENTER)
add_text(slide14, "Why am I doing this? Am I growing, or am I simply staying busy?", Inches(1), Inches(4.5), Inches(11.33), Inches(1), Pt(20), color=TEXT_GRAY, align=PP_ALIGN.CENTER)

# Slide 15: My Cause & Thank You
slide15 = add_blank_slide()
add_text(slide15, "My Cause", Inches(1), Inches(1), Inches(11.33), Inches(1), Pt(40), bold=True, align=PP_ALIGN.CENTER)
add_text(slide15, "To make my parents proud and honor their sacrifices.", Inches(1), Inches(2.5), Inches(11.33), Inches(1), Pt(28), color=TEXT_GRAY, align=PP_ALIGN.CENTER)
add_text(slide15, "THANK YOU", Inches(1), Inches(4.5), Inches(11.33), Inches(1.5), Pt(64), bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)
add_text(slide15, "Pranesh M S | Aspiring Machine Learning Engineer", Inches(1), Inches(6.5), Inches(11.33), Inches(0.5), Pt(14), color=TEXT_GRAY, align=PP_ALIGN.CENTER)

prs.save("About_Me_Pranesh_Premium.pptx")
print("Presentation generated successfully!")
